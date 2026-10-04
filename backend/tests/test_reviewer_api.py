"""Reviewer edits, and the /rate-card, /estimates/reprice and /estimates/export endpoints.

No LLM anywhere: the lines are built by hand, as the mapping step would
leave them, and priced against the Northline Mechanical sample rate card.
"""

from decimal import Decimal as D
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.extraction import expand_production_rate
from app.main import app
from app.pricing import LineEdit, ReviewerEditError, apply_reviewer_edits
from app.schemas import BidLineItem, Flag, LineItemComponent

from tests.test_pipeline import assert_foots

client = TestClient(app)


def flag(severity, code, **extra):
    return Flag(severity=severity, code=code, message=f"{code} message", **extra)


@pytest.fixture
def lines(sample_rate_card) -> list[BidLineItem]:
    """Four draft lines, each needing a different reviewer action.

    L001  thermostats on the company rate PR-TSTAT-PROG, nothing to do
    L002  supply registers on quantities assumed by the LLM, although the
          company has PR-REG-SUP
    L003  refrigerant charge with the quantity written as "TBD"
    L004  a line the LLM could not map at all
    """
    rates = {rate.code: rate for rate in sample_rate_card.production_rates}
    thermostat, thermostat_assumptions = expand_production_rate(
        rates["PR-TSTAT-PROG"], sample_rate_card
    )
    charge, charge_assumptions = expand_production_rate(
        rates["PR-REFRIG-CHARGE"], sample_rate_card
    )
    return [
        BidLineItem(
            id="L001",
            item_number="C1",
            description="Thermostats, programmable",
            quantity=D("6"),
            unit="EA",
            components=thermostat,
            production_rate_code="PR-TSTAT-PROG",
            assumptions=thermostat_assumptions,
            source_ref="p1-t1-r14",
            mapping_confidence="high",
            mapping_rationale="Direct match.",
        ),
        BidLineItem(
            id="L002",
            item_number="B4",
            description="Supply registers",
            quantity=D("64"),
            unit="Each",
            components=[
                LineItemComponent(
                    type="material", rate_card_code="MAT-REG-SUP", quantity_per_unit=D("1")
                ),
                LineItemComponent(
                    type="labor", rate_card_code="LAB-TECH", quantity_per_unit=D("0.5")
                ),
            ],
            flags=[
                flag("warning", "ASSUMED_PRODUCTION_RATE"),
                flag(
                    "warning",
                    "STANDARD_RATE_DECLINED",
                    suggested_production_rate_code="PR-REG-SUP",
                ),
            ],
            assumptions=["Labor LAB-TECH (HVAC Technician): 0.5 hrs per Each [ASSUMED by the LLM]"],
            source_ref="p1-t1-r11",
            mapping_confidence="medium",
            mapping_rationale="No size is stated.",
        ),
        BidLineItem(
            id="L003",
            item_number="B-5",
            description="Refrigerant R-454B, system charge",
            quantity=D("0"),
            unit="LB",
            components=charge,
            production_rate_code="PR-REFRIG-CHARGE",
            flags=[flag("blocker", "QUANTITY_NOT_NUMERIC")],
            assumptions=charge_assumptions,
            source_ref="Base Bid!R9",
        ),
        BidLineItem(
            id="L004",
            item_number="B-6",
            description="Kitchen exhaust hood fire suppression system",
            quantity=D("1"),
            unit="EA",
            flags=[flag("warning", "NO_RATE_CARD_MATCH")],
            source_ref="Base Bid!R10",
        ),
    ]


def by_id(estimate, line_id):
    return next(line for line in estimate.lines if line.id == line_id)


def open_codes(line):
    return [(f.severity, f.code) for f in line.flags if not f.resolved]


# apply_reviewer_edits()


def test_no_edits_reproduces_the_draft(lines, sample_rate_card):
    estimate = apply_reviewer_edits(lines, [], sample_rate_card)

    # L001: 6 x $145.00 = 870.00 material, 6 x 1.25 x $92.00 = 690.00 labor
    # L002: 64 x $18.50 = 1,184.00 material, 64 x 0.5 x $85.00 = 2,720.00 labor
    assert [line.line_subtotal for line in estimate.lines] == [
        D("1560.00"), D("3904.00"), D("0.00"), D("0.00"),
    ]  # fmt: skip
    assert [line.review_status for line in estimate.lines] == ["open"] * 4
    # L002 is far above 15% of the subtotal, so its assumed rate is a blocker.
    assert ("blocker", "ASSUMED_PRODUCTION_RATE") in open_codes(by_id(estimate, "L002"))
    assert ("blocker", "NO_COMPONENTS") in open_codes(by_id(estimate, "L004"))
    assert estimate.review.blocker_count == 3
    assert estimate.review.ready_to_approve is False
    assert estimate.source_lines == lines
    assert_foots(estimate)


def test_quantity_edit_reprices_the_line_and_clears_the_quantity_flag(lines, sample_rate_card):
    estimate = apply_reviewer_edits(
        lines, [LineEdit(line_id="L003", quantity=D("40"))], sample_rate_card
    )

    line = by_id(estimate, "L003")
    # PR-REFRIG-CHARGE: 40 x 1 x $28.00 = 1,120.00 material
    #                   40 x 0.05 x $85.00 = 170.00 labor
    assert (line.material_cost, line.labor_cost) == (D("1120.00"), D("170.00"))
    assert line.quantity == D("40")
    assert line.reviewer_edits == ["quantity"]
    assert "QUANTITY_NOT_NUMERIC" not in [f.code for f in line.flags]
    assert line.assumptions[-1] == "Quantity 40 LB set by the reviewer (the draft had 0)."
    assert_foots(estimate)


def test_applying_the_suggested_rate_replaces_the_llms_numbers(lines, sample_rate_card):
    suggested = lines[1].flags[1].suggested_production_rate_code
    estimate = apply_reviewer_edits(
        lines, [LineEdit(line_id="L002", production_rate_code=suggested)], sample_rate_card
    )

    line = by_id(estimate, "L002")
    # PR-REG-SUP: 64 x 1 x $18.50 = 1,184.00 material
    #             64 x 0.3 x $85.00 = 1,632.00 labor (the LLM had guessed 0.5 hrs)
    assert (line.material_cost, line.labor_cost) == (D("1184.00"), D("1632.00"))
    assert (line.rate_basis, line.production_rate_code) == ("standard", "PR-REG-SUP")
    assert line.reviewer_edits == ["production_rate"]
    # The flags about the LLM's mapping are gone; the line is still big.
    assert open_codes(line) == [("warning", "HIGH_IMPACT_LINE")]
    assert line.assumptions[0] == (
        "Standard production rate PR-REG-SUP (Supply register, 10 x 6, installed), per EA "
        "[chosen by the reviewer; the draft used quantities assumed by the LLM]"
    )
    assert estimate.review.standard_rate_pct == D("100.0")
    assert_foots(estimate)


def test_rate_in_another_unit_is_a_unit_mismatch(lines, sample_rate_card):
    estimate = apply_reviewer_edits(
        lines, [LineEdit(line_id="L004", production_rate_code="PR-DUCT-FLEX-8")], sample_rate_card
    )

    line = by_id(estimate, "L004")
    assert line.line_subtotal == D("0.00")
    assert ("blocker", "UNIT_MISMATCH") in open_codes(line)


def test_reviewed_resolves_warnings_and_confirmable_blockers_only(lines, sample_rate_card):
    estimate = apply_reviewer_edits(
        lines,
        [
            LineEdit(line_id="L002", status="reviewed"),
            LineEdit(line_id="L003", status="reviewed"),
        ],
        sample_rate_card,
    )

    # Reviewing confirms the assumed rate on L002 ...
    assumed = by_id(estimate, "L002")
    assert assumed.review_status == "reviewed"
    assert open_codes(assumed) == []
    assert all(f.resolved for f in assumed.flags)
    # ... but cannot make up for the missing quantity on L003.
    assert open_codes(by_id(estimate, "L003")) == [("blocker", "QUANTITY_NOT_NUMERIC")]
    # Prices do not change when a line is only marked reviewed.
    assert assumed.line_subtotal == D("3904.00")
    # Left: QUANTITY_NOT_NUMERIC on L003 and NO_COMPONENTS on L004.
    assert estimate.review.blocker_count == 2
    assert estimate.review.ready_to_approve is False


def test_excluded_line_adds_nothing_and_no_longer_blocks(lines, sample_rate_card):
    estimate = apply_reviewer_edits(
        lines,
        [LineEdit(line_id="L002", status="excluded", exclusion_reason="  By owner. ")],
        sample_rate_card,
    )

    line = by_id(estimate, "L002")
    assert line.line_subtotal == D("0")
    assert (line.review_status, line.exclusion_reason) == ("excluded", "By owner.")
    assert line.calculation_trace == ["Not priced: excluded by the reviewer (By owner.)"]
    assert line.subtotal_share_pct is None
    assert all(f.resolved for f in line.flags)
    # Only L001 is priced now, so it is 100% of the subtotal.
    assert by_id(estimate, "L001").subtotal_share_pct == D("100.0")
    assert estimate.totals.direct_labor == D("690.00")
    assert estimate.review.lines_excluded == 1
    assert [line.id for line in estimate.lines] == ["L001", "L002", "L003", "L004"]
    assert_foots(estimate)


def test_ready_to_approve_once_every_blocker_is_resolved(lines, sample_rate_card):
    estimate = apply_reviewer_edits(
        lines,
        [
            LineEdit(line_id="L002", production_rate_code="PR-REG-SUP", status="reviewed"),
            LineEdit(line_id="L003", quantity=D("40")),
            LineEdit(line_id="L004", status="excluded", exclusion_reason="Subcontracted."),
        ],
        sample_rate_card,
    )

    assert estimate.review.blocker_count == 0
    assert estimate.review.ready_to_approve is True
    assert_foots(estimate)


def test_nothing_is_ready_to_approve_when_every_line_is_excluded(lines, sample_rate_card):
    estimate = apply_reviewer_edits(
        lines,
        [LineEdit(line_id=line.id, status="excluded", exclusion_reason="Out.") for line in lines],
        sample_rate_card,
    )

    assert estimate.totals.grand_total == D("0")
    assert estimate.review.blocker_count == 0
    assert estimate.review.ready_to_approve is False


def test_edits_do_not_modify_the_draft_lines(lines, sample_rate_card):
    before = [line.model_copy(deep=True) for line in lines]

    apply_reviewer_edits(
        lines,
        [LineEdit(line_id="L002", production_rate_code="PR-REG-SUP", quantity=D("10"))],
        sample_rate_card,
    )

    assert lines == before


@pytest.mark.parametrize(
    ("edit", "message"),
    [
        (LineEdit(line_id="L999", quantity=D("1")), "unknown line 'L999'"),
        (LineEdit(line_id="L001", production_rate_code="PR-NOPE"), "'PR-NOPE' is not in"),
        (LineEdit(line_id="L001", status="excluded"), "a reason is required"),
        (LineEdit(line_id="L001", status="excluded", exclusion_reason="   "), "a reason is required"),
    ],
)  # fmt: skip
def test_invalid_edits_are_rejected(lines, sample_rate_card, edit, message):
    with pytest.raises(ReviewerEditError, match=message):
        apply_reviewer_edits(lines, [edit], sample_rate_card)


def test_two_edits_for_one_line_are_rejected(lines, sample_rate_card):
    edits = [LineEdit(line_id="L001", quantity=D("1")), LineEdit(line_id="L001", quantity=D("2"))]
    with pytest.raises(ReviewerEditError, match="More than one edit"):
        apply_reviewer_edits(lines, edits, sample_rate_card)


# GET /rate-card


def test_rate_card_endpoint_returns_the_card_with_its_production_rates():
    response = client.get("/rate-card")

    assert response.status_code == 200
    body = response.json()
    assert (body["company_name"], body["trade"], body["currency"]) == (
        "Northline Mechanical", "HVAC", "USD",
    )  # fmt: skip
    assert [len(body[key]) for key in ("labor_rates", "materials", "equipment")] == [5, 15, 4]
    assert len(body["production_rates"]) == 15
    register = next(r for r in body["production_rates"] if r["code"] == "PR-REG-SUP")
    assert (register["unit"], register["components"][1]["quantity_per_unit"]) == ("EA", "0.3")


def test_samples_endpoint_lists_the_three_sample_bids():
    response = client.get("/samples")

    assert response.status_code == 200
    assert [sample["id"] for sample in response.json()] == ["clean", "messy", "tricky"]


def test_unknown_sample_is_a_404():
    # 404 is raised before the pipeline runs, so no LLM is involved.
    from app.main import llm_client

    app.dependency_overrides[llm_client] = lambda: None
    try:
        response = client.post("/estimates/draft/sample/nope")
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 404


def test_cors_allows_the_frontend_origin():
    response = client.get("/rate-card", headers={"Origin": "http://localhost:3000"})

    assert response.headers["access-control-allow-origin"] == "http://localhost:3000"


# POST /estimates/reprice


def payload(lines, edits=(), **extra):
    return {
        "lines": [line.model_dump(mode="json") for line in lines],
        "edits": list(edits),
        **extra,
    }


def test_reprice_endpoint_returns_a_fresh_estimate_and_review(lines):
    skipped = [{"source_ref": "p1-t1-r6", "description": "Subtotal", "reason": "subtotal / total row"}]
    response = client.post(
        "/estimates/reprice",
        json=payload(
            lines,
            [
                {"line_id": "L002", "production_rate_code": "PR-REG-SUP", "status": "reviewed"},
                {"line_id": "L003", "quantity": "40"},
                {"line_id": "L004", "status": "excluded", "exclusion_reason": "Subcontracted."},
            ],
            skipped_rows=skipped,
        ),
    )

    assert response.status_code == 200
    body = response.json()
    assert [line["line_subtotal"] for line in body["lines"]] == [
        "1560.00", "2816.00", "1290.00", "0",
    ]  # fmt: skip
    assert [line["review_status"] for line in body["lines"]] == [
        "open", "reviewed", "open", "excluded",
    ]  # fmt: skip
    # labor 690.00 + 1,632.00 + 170.00; material 870.00 + 1,184.00 + 1,120.00
    assert body["totals"]["direct_labor"] == "2492.00"
    assert body["totals"]["direct_material"] == "3174.00"
    assert body["review"]["blocker_count"] == 0
    assert body["review"]["ready_to_approve"] is True
    assert body["review"]["standard_rate_pct"] == "100.0"
    assert body["skipped_rows"] == skipped
    # The draft lines come back untouched, ready for the next reprice.
    assert body["source_lines"][2]["quantity"] == "0"


def test_reprice_is_repeatable(lines):
    request = payload(lines, [{"line_id": "L003", "quantity": "40"}])

    first = client.post("/estimates/reprice", json=request).json()
    second = client.post("/estimates/reprice", json=request).json()

    assert first == second


def test_reprice_rejects_an_exclusion_without_a_reason(lines):
    response = client.post(
        "/estimates/reprice", json=payload(lines, [{"line_id": "L001", "status": "excluded"}])
    )

    assert response.status_code == 422
    assert "a reason is required" in response.json()["detail"]


def test_reprice_rejects_an_unknown_production_rate(lines):
    response = client.post(
        "/estimates/reprice",
        json=payload(lines, [{"line_id": "L001", "production_rate_code": "PR-NOPE"}]),
    )

    assert response.status_code == 422
    assert "PR-NOPE" in response.json()["detail"]


# POST /estimates/export


AUDIT_LOG = [
    {
        "timestamp": "2026-10-04T10:15:00.000Z",
        "who": "Estimator",
        "action": "Changed quantity",
        "line_id": "L003",
        "item_number": "B-5",
        "before": "0 LB",
        "after": "40 LB",
        "note": "Line subtotal $0.00 -> $1,290.00",
    }
]


def export(lines, edits=(), audit_log=AUDIT_LOG):
    response = client.post(
        "/estimates/export",
        json=payload(lines, edits, audit_log=audit_log, source_filename="tricky bid.xlsx"),
    )
    assert response.status_code == 200
    return response, load_workbook(BytesIO(response.content))


def rows(sheet):
    return [list(row) for row in sheet.iter_rows(values_only=True)]


def test_export_returns_an_xlsx_with_the_three_sheets(lines):
    response, workbook = export(lines, [{"line_id": "L003", "quantity": "40"}])

    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert response.headers["content-disposition"] == (
        'attachment; filename="tricky_bid_estimate.xlsx"'
    )
    assert workbook.sheetnames == ["Estimate", "Flags and assumptions", "Audit log"]


def test_export_estimate_sheet_matches_the_repriced_estimate(lines, sample_rate_card):
    edits = [
        {"line_id": "L003", "quantity": "40"},
        {"line_id": "L004", "status": "excluded", "exclusion_reason": "Subcontracted."},
    ]
    _, workbook = export(lines, edits)
    estimate = apply_reviewer_edits(
        lines, [LineEdit.model_validate(edit) for edit in edits], sample_rate_card
    )

    sheet = rows(workbook["Estimate"])
    labelled = {row[0]: row[1] for row in sheet if row[0] and row[1] is not None}
    assert labelled["Company"] == "Northline Mechanical"
    assert labelled["Bid schedule"] == "tricky bid.xlsx"
    # L002's assumed rate is still an unresolved blocker.
    assert labelled["Status"] == "NOT ready to approve: 1 unresolved blocker(s)"

    header = next(index for index, row in enumerate(sheet) if row[0] == "Line")
    line_rows = sheet[header + 1 : header + 5]
    assert [row[0] for row in line_rows] == ["L001", "L002", "L003", "L004"]
    assert [D(str(row[10])) for row in line_rows] == [
        line.line_subtotal for line in estimate.lines
    ]
    assert [row[12] for row in line_rows] == ["Open", "Open", "Open", "Excluded"]
    assert line_rows[3][14] == "Subcontracted."

    totals = {row[9]: D(str(row[10])) for row in sheet[header + 5 :] if row[9]}
    assert totals == {
        "Direct labor": estimate.totals.direct_labor,
        "Direct material": estimate.totals.direct_material,
        "Material markup": estimate.totals.material_markup,
        "Sales tax on materials": estimate.totals.sales_tax,
        "Direct equipment": estimate.totals.direct_equipment,
        "Subtotal": estimate.totals.subtotal,
        "Overhead": estimate.totals.overhead,
        "Profit": estimate.totals.profit,
        "Grand total": estimate.totals.grand_total,
    }
    # The file carries values only: no formulas to recalculate.
    assert not [cell for row in sheet for cell in row if isinstance(cell, str) and cell.startswith("=")]


def test_export_lists_flags_and_assumptions_with_their_status(lines):
    _, workbook = export(lines, [{"line_id": "L002", "status": "reviewed"}])

    sheet = rows(workbook["Flags and assumptions"])
    assert sheet[0] == ["Line", "Item", "Description", "Type", "Severity", "Code", "Status", "Detail"]
    flags = {(row[0], row[5]): row[6] for row in sheet[1:] if row[3] == "Flag"}
    assert flags[("L002", "ASSUMED_PRODUCTION_RATE")] == "Resolved"
    assert flags[("L003", "QUANTITY_NOT_NUMERIC")] == "Open"
    assumptions = [(row[0], row[6]) for row in sheet[1:] if row[3] == "Assumption"]
    assert ("L001", "Open") in assumptions
    assert ("L002", "Accepted") in assumptions


def test_export_audit_log_sheet(lines):
    _, workbook = export(lines)

    assert rows(workbook["Audit log"]) == [
        ["Time", "Who", "Action", "Line", "Item", "Before", "After", "Note"],
        [
            "2026-10-04T10:15:00.000Z", "Estimator", "Changed quantity", "L003", "B-5",
            "0 LB", "40 LB", "Line subtotal $0.00 -> $1,290.00",
        ],
    ]  # fmt: skip


def test_export_with_an_empty_audit_log(lines):
    _, workbook = export(lines, audit_log=[])

    assert rows(workbook["Audit log"])[1][0] == "No reviewer actions were recorded."


def test_export_rejects_invalid_edits(lines):
    response = client.post(
        "/estimates/export", json=payload(lines, [{"line_id": "L999", "quantity": "1"}])
    )

    assert response.status_code == 422
