"""Excel export of a reviewed estimate.

Three sheets: the estimate, its flags and assumptions, and the reviewer audit
log. Every figure is copied from an Estimate the engine has just priced;
nothing is calculated here and the workbook contains no formulas.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel

from app.schemas import Estimate, RateCard

MONEY_FORMAT = "#,##0.00"
BOLD = Font(bold=True)

_STATUS = {"open": "Open", "reviewed": "Reviewed", "excluded": "Excluded"}
_BASIS = {"standard": "Company standard", "assumed": "Assumed by the LLM", "none": "Not priced"}


class AuditEntry(BaseModel):
    """One reviewer action, as recorded by the frontend."""

    timestamp: str
    who: str
    action: str
    line_id: str | None = None
    item_number: str | None = None
    before: str | None = None
    after: str | None = None
    note: str | None = None


def build_workbook(
    estimate: Estimate,
    audit_log: list[AuditEntry],
    rate_card: RateCard,
    source_filename: str | None = None,
) -> bytes:
    workbook = Workbook()
    _estimate_sheet(workbook.active, estimate, rate_card, source_filename)
    _flags_sheet(workbook.create_sheet(), estimate)
    _audit_sheet(workbook.create_sheet(), audit_log)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _estimate_sheet(
    sheet: Worksheet, estimate: Estimate, rate_card: RateCard, source_filename: str | None
) -> None:
    sheet.title = "Estimate"
    review = estimate.review
    if review.ready_to_approve:
        status = "Ready to approve: no unresolved blockers"
    else:
        status = f"NOT ready to approve: {review.blocker_count} unresolved blocker(s)"

    sheet.append(["Draft estimate"])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append(["Company", rate_card.company_name])
    sheet.append(["Trade", rate_card.trade])
    sheet.append(["Bid schedule", source_filename or ""])
    sheet.append(["Currency", estimate.currency])
    sheet.append(["Status", status])
    sheet.append(["Subtotal on company standard rates (%)", review.standard_rate_pct])
    sheet.append(["Subtotal on rates assumed by the LLM (%)", review.assumed_rate_pct])
    sheet.append([])

    _header(
        sheet,
        [
            "Line", "Item", "Description", "Qty", "Unit", "Production rate", "Rate basis",
            "Labor", "Material", "Equipment", "Line subtotal", "% of subtotal",
            "Review status", "Open flags", "Exclusion reason", "Source",
        ],
    )  # fmt: skip
    for line in estimate.lines:
        open_flags = [flag.code for flag in line.flags if not flag.resolved]
        sheet.append(
            [
                line.id,
                line.item_number,
                line.description,
                line.quantity,
                line.unit,
                line.production_rate_code or "",
                _BASIS[line.rate_basis],
                line.labor_cost,
                line.material_cost,
                line.equipment_cost,
                line.line_subtotal,
                line.subtotal_share_pct,
                _STATUS[line.review_status],
                ", ".join(open_flags),
                line.exclusion_reason or "",
                line.source_ref or "",
            ]
        )
        for column in "HIJK":
            sheet[f"{column}{sheet.max_row}"].number_format = MONEY_FORMAT

    sheet.append([])
    totals = estimate.totals
    for label, amount in (
        ("Direct labor", totals.direct_labor),
        ("Direct material", totals.direct_material),
        ("Material markup", totals.material_markup),
        ("Sales tax on materials", totals.sales_tax),
        ("Direct equipment", totals.direct_equipment),
        ("Subtotal", totals.subtotal),
        ("Overhead", totals.overhead),
        ("Profit", totals.profit),
        ("Grand total", totals.grand_total),
    ):
        sheet.append([None] * 9 + [label, amount])
        sheet[f"K{sheet.max_row}"].number_format = MONEY_FORMAT
        if label in ("Subtotal", "Grand total"):
            sheet[f"J{sheet.max_row}"].font = BOLD
            sheet[f"K{sheet.max_row}"].font = BOLD

    _widths(sheet, [8, 10, 52, 10, 8, 22, 20, 14, 14, 22, 16, 14, 14, 44, 30, 18])


def _flags_sheet(sheet: Worksheet, estimate: Estimate) -> None:
    """Every flag, then the assumptions behind each line.

    An assumption is "Open" until its line is reviewed, "Accepted" once it
    is, and "Excluded" when the line was taken out of the estimate.
    """
    sheet.title = "Flags and assumptions"
    _header(sheet, ["Line", "Item", "Description", "Type", "Severity", "Code", "Status", "Detail"])
    for line in estimate.lines:
        for flag in line.flags:
            sheet.append(
                [
                    line.id,
                    line.item_number,
                    line.description,
                    "Flag",
                    flag.severity,
                    flag.code,
                    "Resolved" if flag.resolved else "Open",
                    flag.message,
                ]
            )
    assumption_status = {"open": "Open", "reviewed": "Accepted", "excluded": "Excluded"}
    for line in estimate.lines:
        for assumption in line.assumptions:
            sheet.append(
                [
                    line.id,
                    line.item_number,
                    line.description,
                    "Assumption",
                    "",
                    "",
                    assumption_status[line.review_status],
                    assumption,
                ]
            )
    for row in sheet.iter_rows(min_row=2, min_col=8, max_col=8):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    _widths(sheet, [8, 10, 44, 12, 10, 28, 10, 110])


def _audit_sheet(sheet: Worksheet, audit_log: list[AuditEntry]) -> None:
    sheet.title = "Audit log"
    _header(sheet, ["Time", "Who", "Action", "Line", "Item", "Before", "After", "Note"])
    for entry in audit_log:
        sheet.append(
            [
                entry.timestamp,
                entry.who,
                entry.action,
                entry.line_id or "",
                entry.item_number or "",
                entry.before or "",
                entry.after or "",
                entry.note or "",
            ]
        )
    if not audit_log:
        sheet.append(["No reviewer actions were recorded."])
    _widths(sheet, [26, 14, 30, 8, 10, 36, 36, 50])


def _header(sheet: Worksheet, titles: list[str]) -> None:
    sheet.append(titles)
    for cell in sheet[sheet.max_row]:
        cell.font = BOLD
    # A coordinate, not sheet.cell(): that would create the row and push the data down.
    sheet.freeze_panes = f"A{sheet.max_row + 1}"


def _widths(sheet: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths):
        sheet.column_dimensions[chr(ord("A") + index)].width = width
