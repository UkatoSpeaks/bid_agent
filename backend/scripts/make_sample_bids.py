"""Regenerate the synthetic sample bid schedules in data/bids/.

    uv run python scripts/make_sample_bids.py

All three are fictional HVAC bid schedules for Northline Mechanical, written
to fit data/rate_cards/hvac_rate_card.json:

    clean_bid.xlsx   a tidy table of 10 items
    messy_bid.pdf    header, section titles, a subtotal row, inconsistent
                     units and one vague item
    tricky_bid.xlsx  two sheets, one item with no rate card match and one
                     quantity written as "TBD"
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

BIDS_DIR = Path(__file__).resolve().parent.parent / "data" / "bids"

HEADER = ["Item No.", "Description", "Qty", "Unit"]


def make_clean_bid(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bid Schedule"
    rows = [
        ["1", "Furnish and install 3-ton condensing unit", 4, "EA"],
        ["2", "Furnish and install 3-ton air handler", 4, "EA"],
        ["3", "Furnish and install 5-ton condensing unit", 2, "EA"],
        ["4", "Furnish and install 5-ton air handler", 2, "EA"],
        ["5", "Refrigerant line set, 3/8 x 7/8, 50 ft, insulated", 6, "EA"],
        ["6", "12 in. spiral round duct, installed", 480, "LF"],
        ["7", "8 in. insulated flexible duct, installed", 350, "LF"],
        ["8", "Programmable thermostat, installed and wired", 6, "EA"],
        ["9", "Supply register, 10 x 6", 48, "EA"],
        ["10", "Return air grille, 20 x 20", 12, "EA"],
    ]
    _write_rows(sheet, [HEADER, *rows], header_row=1)
    workbook.save(path)


def make_tricky_bid(path: Path) -> None:
    workbook = Workbook()

    base = workbook.active
    base.title = "Base Bid"
    base_rows = [
        ["Northline Mechanical - Bid Schedule", None, None, None],
        ["Project: Harbor Street Clinic, HVAC Renovation", None, None, None],
        [None, None, None, None],
        HEADER,
        ["B-1", "Packaged rooftop unit, 10 ton, including crane set", 2, "EA"],
        ["B-2", "Galvanized sheet metal ductwork, fabricated and installed", 3200, "LB"],
        ["B-3", "Duct wrap insulation, R-6 foil-faced", 1850.5, "SF"],
        ["B-4", "Volume control damper, 12 in", 14, "EA"],
        ["B-5", "Refrigerant R-454B, system charge", "TBD", "LB"],
        ["B-6", "Kitchen exhaust hood fire suppression system", 1, "EA"],
    ]
    _write_rows(base, base_rows, header_row=4)

    alternates = workbook.create_sheet("Alternates")
    alternate_rows = [
        ["Alternates (priced separately)", None, None, None],
        HEADER,
        ["ALT-1", "Add programmable thermostat at each exam room", 9, "EA"],
        ["ALT-2", "Add return air grille, 20 x 20, at corridor", 6, "EA"],
        ["ALT-3", "Recover refrigerant from existing rooftop units", 3, "EA"],
    ]
    _write_rows(alternates, alternate_rows, header_row=2)

    workbook.save(path)


def _write_rows(sheet, rows: list[list], header_row: int) -> None:
    for row in rows:
        sheet.append(row)
    for cell in sheet[header_row]:
        cell.font = Font(bold=True)
    sheet.column_dimensions["A"].width = 10
    sheet.column_dimensions["B"].width = 60


def make_messy_bid(path: Path) -> None:
    styles = getSampleStyleSheet()
    story = [
        Paragraph("NORTHLINE MECHANICAL", styles["Title"]),
        Paragraph("Bid Schedule - Lakeview Elementary School, HVAC Replacement", styles["Heading3"]),
        Paragraph("Bid No. 2026-114 &nbsp;&nbsp; Bids due: October 30, 2026", styles["Normal"]),
        Spacer(1, 0.2 * inch),
    ]

    rows = [
        ["ITEM", "DESCRIPTION OF WORK", "QTY", "UNIT"],
        ["SECTION A - EQUIPMENT", "", "", ""],
        ["A1", "Packaged rooftop unit, 10 ton, set by crane", "3", "EA"],
        ["A2", "5-ton split system: condensing unit", "2", "each"],
        ["A3", "5-ton split system: air handler", "2", "each"],
        ["A4", "Refrigerant line sets, 50 ft", "2", "ea."],
        ["", "Subtotal Section A", "9", ""],
        ["SECTION B - AIR DISTRIBUTION", "", "", ""],
        ["B1", 'Spiral duct, 12" round', "1,150", "LF"],
        ["B2", '8" flex duct to diffusers', "275", "lin. ft"],
        ["B3", "Duct insulation wrap (R-6)", "2,400", "SF"],
        ["B4", "Supply registers", "64", "Each"],
        ["B5", "Misc. ductwork modifications as required", "1", "LS"],
        ["SECTION C - CONTROLS", "", "", ""],
        ["C1", "Thermostats, programmable", "11", "EA"],
    ]
    section_rows = [i for i, row in enumerate(rows) if row[0].startswith("SECTION")]
    table = Table(rows, colWidths=[0.7 * inch, 4.2 * inch, 0.8 * inch, 0.8 * inch])
    style = [
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("FONTNAME", (1, 6), (2, 6), "Helvetica-Oblique"),
    ]
    for index in section_rows:
        style.append(("SPAN", (0, index), (-1, index)))
        style.append(("FONTNAME", (0, index), (-1, index), "Helvetica-Bold"))
    table.setStyle(TableStyle(style))
    story.append(table)
    story.append(Spacer(1, 0.2 * inch))
    story.append(
        Paragraph(
            "Notes: Quantities are approximate. Bidder to field verify all "
            "dimensions before fabrication.",
            styles["Normal"],
        )
    )

    # invariant=1 keeps the file byte-identical between runs (no timestamps).
    SimpleDocTemplate(str(path), pagesize=letter, invariant=1, title="Bid Schedule").build(story)


def main() -> None:
    BIDS_DIR.mkdir(parents=True, exist_ok=True)
    make_clean_bid(BIDS_DIR / "clean_bid.xlsx")
    make_messy_bid(BIDS_DIR / "messy_bid.pdf")
    make_tricky_bid(BIDS_DIR / "tricky_bid.xlsx")
    for name in ("clean_bid.xlsx", "messy_bid.pdf", "tricky_bid.xlsx"):
        print(f"wrote {BIDS_DIR / name}")


if __name__ == "__main__":
    main()
